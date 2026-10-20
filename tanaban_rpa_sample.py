# -*- coding: utf-8 -*-
"""
棚番マスタ自動更新RPAツール（ポートフォリオ用サンプル）

本サンプルでは、操作用GUI（Tkinter）はそのまま掲載しています。

一方、基幹システムを自動操作するRPA処理（pyautogui / pywinauto）は、
機密情報保護のためログ出力によるエミュレーションへ置き換えています。
"""

import os
import time
import datetime
import threading
import tkinter as tk
from tkinter import messagebox
import openpyxl


class TanabanRpaApp:
    def __init__(self, root):
        self.root = root
        self.root.title("棚番マスタ更新ツール (Sample Ver)")
        self.root.geometry("400x200")

        self.lbl_code = tk.Label(
            root,
            text="倉庫コードを入力してください:",
            font=("Meiryo UI", 11),
        )
        self.lbl_code.pack(pady=10)

        self.entry_code = tk.Entry(
            root,
            font=("Meiryo UI", 11),
            width=15,
        )
        self.entry_code.pack(pady=5)
        self.entry_code.focus_set()

        self.btn_frame = tk.Frame(root)
        self.btn_frame.pack(pady=20)

        self.btn_start = tk.Button(
            self.btn_frame,
            text="スタート",
            font=("Meiryo UI", 10),
            width=10,
            command=self.start_process,
        )
        self.btn_start.pack(side=tk.LEFT, padx=10)

        self.btn_exit = tk.Button(
            self.btn_frame,
            text="終了",
            font=("Meiryo UI", 10),
            width=10,
            command=root.quit,
        )
        self.btn_exit.pack(side=tk.LEFT, padx=10)

        self.worker_thread = None

    def start_process(self):
        code = self.entry_code.get().strip()
        if not code:
            messagebox.showwarning(
                "入力エラー",
                "倉庫コードを入力してください。",
            )
            return

        self.btn_start.config(state=tk.DISABLED)

        self.worker_thread = threading.Thread(
            target=self.rpa_core_process,
            args=(code,),
        )
        self.worker_thread.daemon = True
        self.worker_thread.start()

    def rpa_core_process(self, warehouse_code):
        excel_filename = "input_sample.xlsx"

        if not os.path.exists(excel_filename):
            messagebox.showerror(
                "ファイルなし",
                f"入力ファイル '{excel_filename}' が見つかりません。",
            )
            self.root.after(
                0,
                lambda: self.btn_start.config(state=tk.NORMAL),
            )
            return

        print(f"=== 処理開始 (倉庫コード: {warehouse_code}) ===")

        try:
            wb = openpyxl.load_workbook(excel_filename)
            ws = wb.active

            for row in range(2, ws.max_row + 1):
                item_code = ws.cell(row=row, column=1).value
                new_tanaban = ws.cell(row=row, column=2).value

                if not item_code:
                    continue

                print(
                    f"[{row-1}件目] 品目: {item_code} "
                    f"-> 新棚番: {new_tanaban} "
                    "を基幹システムに入力中..."
                )

                time.sleep(0.5)

                if "ERROR" in str(item_code):
                    print(
                        "【エラー検知】 "
                        "基幹システム側でエラーダイアログが表示されました。"
                    )
                    ws.cell(row=row, column=3).value = "失敗"
                    ws.cell(row=row, column=4).value = (
                        "エラー: 該当データなし"
                    )
                    ws.cell(row=row, column=5).value = (
                        datetime.datetime.now().strftime(
                            "%Y-%m-%d %H:%M:%S"
                        )
                    )
                    wb.save(excel_filename)

                    messagebox.showerror(
                        "処理失敗",
                        f"品目 {item_code} の更新中にエラーを検出したため、"
                        "処理を中断します。",
                    )
                    return

                ws.cell(row=row, column=3).value = "成功"
                ws.cell(row=row, column=4).value = "Batch 1 成功"
                ws.cell(row=row, column=5).value = (
                    datetime.datetime.now().strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )
                )

            wb.save(excel_filename)
            messagebox.showinfo(
                "正常終了",
                "すべての棚番更新処理が正常に完了しました。",
            )

        except Exception as e:
            messagebox.showerror(
                "致命的エラー",
                f"処理中に予期せぬエラーが発生しました:\n{str(e)}",
            )

        finally:
            self.root.after(
                0,
                lambda: self.btn_start.config(state=tk.NORMAL),
            )


if __name__ == "__main__":
    root = tk.Tk()
    app = TanabanRpaApp(root)
    root.mainloop()